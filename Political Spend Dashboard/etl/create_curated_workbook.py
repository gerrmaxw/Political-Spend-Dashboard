from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path

import pandas as pd
from pandas.errors import EmptyDataError
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

from config import BASE_DIR, CURATED_DIR, QA_DIR, REFRESH_DATE


OUTPUT_DIR = BASE_DIR / "outputs"
OUTPUT_FILE = OUTPUT_DIR / "PoliticalSpendDashboard.xlsx"

CORE_TABLES = [
    "fact_spend_current",
    "fact_spend_snapshot",
    "fact_outside_spend",
    "fact_political_windows",
    "fact_candidate_finance",
    "dim_snapshot",
    "dim_advertiser",
    "dim_agency",
    "dim_media_type",
    "dim_state",
    "dim_dma",
    "dim_office",
    "dim_party",
    "dim_race_level",
    "dim_station",
    "dim_network",
    "dim_candidate",
    "dim_race",
    "dim_committee",
    "bridge_advertiser_entity",
    "bridge_race_dma",
    "fact_election_results",
    "dim_case_study_notes",
    "dim_data_source",
    "fact_refresh_log",
    "qa_source_health",
    "dim_civic_election",
    "fact_civic_contest",
    "dim_civic_candidate",
]

QA_TABLES = [
    "qa_snapshot_reconciliation",
    "advertiser_match_review",
    "low_confidence_matches",
    "unmatched_advertisers",
    "duplicate_key_check",
    "missing_field_checks",
    "market_scope_exclusions",
]

EMPTY_TABLE_COLUMNS = {
    "fact_outside_spend": [
        "OutsideSpendKey",
        "SourceType",
        "CommitteeKey",
        "CommitteeName",
        "CandidateKey",
        "CandidateName",
        "RaceKey",
        "Office",
        "OfficeCode",
        "StateKey",
        "District",
        "ElectionType",
        "SpendDate",
        "Amount",
        "SupportOppose",
        "Purpose",
        "Payee",
        "FileNumber",
        "TransactionID",
        "ImageNumber",
    ],
}


def safe_sheet_name(name: str, used: set[str]) -> str:
    words = name.replace("fact_", "Fact_").replace("dim_", "Dim_").replace("qa_", "QA_").split("_")
    sheet = "".join(part[:1].upper() + part[1:] for part in words if part)
    sheet = sheet.replace("AdvertiserEntity", "AdvEntity")
    sheet = sheet.replace("PoliticalWindows", "PolWindows")
    sheet = sheet.replace("IndependentExpenditures", "IE")
    sheet = sheet[:31] or "Sheet"
    candidate = sheet
    suffix = 1
    while candidate in used:
        suffix += 1
        candidate = f"{sheet[:28]}{suffix}"
    used.add(candidate)
    return candidate


def prefixed_sheet_name(name: str, used: set[str]) -> str:
    sheet = safe_sheet_name(name, set())
    candidate = f"WS_{sheet}"[:31]
    base = candidate
    suffix = 1
    while candidate in used:
        suffix += 1
        candidate = f"{base[:28]}{suffix}"
    used.add(candidate)
    return candidate


def safe_table_name(name: str, used: set[str]) -> str:
    sheet_style = safe_sheet_name(name, set())
    candidate = "".join(ch if ch.isalnum() or ch == "_" else "_" for ch in sheet_style)
    if not candidate or not candidate[0].isalpha():
        candidate = f"T_{candidate}"
    candidate = candidate[:240]
    base = candidate
    suffix = 1
    while candidate in used:
        suffix += 1
        candidate = f"{base[:236]}_{suffix}"
    used.add(candidate)
    return candidate


def csv_row_count(path: Path) -> int:
    if not path.exists():
        return 0
    with path.open("r", encoding="utf-8", newline="") as handle:
        reader = csv.reader(handle)
        rows = sum(1 for _ in reader)
    return max(rows - 1, 0)


def read_csv(path: Path, columns: list[str] | None = None) -> pd.DataFrame:
    if not path.exists():
        return pd.DataFrame(columns=columns or [])
    try:
        return pd.read_csv(path, dtype=str, keep_default_na=False, low_memory=False)
    except EmptyDataError:
        return pd.DataFrame(columns=columns or [])


def add_dataframe_sheet(
    wb: Workbook,
    name: str,
    df: pd.DataFrame,
    used_sheets: set[str],
    used_tables: set[str],
    *,
    prefix_sheet_names: bool = False,
) -> None:
    if prefix_sheet_names:
        sheet_name = prefixed_sheet_name(name, used_sheets)
    else:
        sheet_name = safe_sheet_name(name, used_sheets)
    ws = wb.create_sheet(sheet_name)
    ws.append(list(df.columns))
    for row in df.itertuples(index=False, name=None):
        ws.append(list(row))

    header_fill = PatternFill("solid", fgColor="8C1D18")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font

    ws.freeze_panes = "A2"
    for col_idx, column_name in enumerate(df.columns, start=1):
        sample_values = [str(column_name)]
        if not df.empty:
            sample_values.extend(str(value) for value in df.iloc[:200, col_idx - 1].tolist())
        width = min(max(max(len(value) for value in sample_values) + 2, 10), 42)
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    if df.columns.size:
        table_name = safe_table_name(name, used_tables)
        table_ref = f"A1:{get_column_letter(df.columns.size)}{max(ws.max_row, 1)}"
        table = Table(displayName=table_name, ref=table_ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)


def add_summary_sheet(wb: Workbook, table_paths: list[tuple[str, Path]]) -> None:
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Political Spend Curated Model"
    ws["A2"] = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws["A3"] = f"Project folder: {BASE_DIR}"
    ws["A5"] = "Power BI should connect to the curated CSV files for refresh. This workbook is a packaged review/upload copy."
    ws["A1"].font = Font(size=16, bold=True, color="8C1D18")
    ws["A5"].font = Font(italic=True)

    row = 7
    ws.cell(row=row, column=1, value="Table")
    ws.cell(row=row, column=2, value="Rows")
    ws.cell(row=row, column=3, value="CSV Source")
    for cell in ws[row]:
        cell.fill = PatternFill("solid", fgColor="8C1D18")
        cell.font = Font(color="FFFFFF", bold=True)

    for table_name, path in table_paths:
        row += 1
        ws.cell(row=row, column=1, value=table_name)
        ws.cell(row=row, column=2, value=csv_row_count(path))
        ws.cell(row=row, column=3, value=str(path))

    reconciliation = read_csv(CURATED_DIR / "qa_snapshot_reconciliation.csv")
    if not reconciliation.empty:
        row += 3
        ws.cell(row=row, column=1, value="Snapshot QA")
        ws.cell(row=row, column=1).font = Font(bold=True, color="8C1D18")
        row += 1
        keep_cols = [
            "SnapshotDate",
            "SourceFile",
            "RowsLoaded",
            "RowsWithAmount",
            "TotalAmount",
            "BroadcastSpend",
            "CableSpend",
            "CTVSpend",
            "DigitalSpend",
            "RadioSpend",
            "TotalAmountDelta",
        ]
        keep_cols = [col for col in keep_cols if col in reconciliation.columns]
        for col_idx, col in enumerate(keep_cols, start=1):
            ws.cell(row=row, column=col_idx, value=col)
            ws.cell(row=row, column=col_idx).fill = PatternFill("solid", fgColor="D9EAD3")
            ws.cell(row=row, column=col_idx).font = Font(bold=True)
        for values in reconciliation[keep_cols].itertuples(index=False, name=None):
            row += 1
            for col_idx, value in enumerate(values, start=1):
                ws.cell(row=row, column=col_idx, value=value)

    ws.freeze_panes = "A8"
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 90


def build_workbook(*, prefix_sheet_names: bool = False) -> Workbook:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    used_sheets = {"Summary"}
    used_tables: set[str] = set()

    table_paths: list[tuple[str, Path]] = []
    for table in CORE_TABLES:
        table_paths.append((table, CURATED_DIR / f"{table}.csv"))
    for table in QA_TABLES:
        table_paths.append((f"qa_{table}" if not table.startswith("qa_") else table, QA_DIR / f"{table}.csv"))

    add_summary_sheet(wb, table_paths)

    for table_name, path in table_paths:
        df = read_csv(path, EMPTY_TABLE_COLUMNS.get(table_name))
        add_dataframe_sheet(wb, table_name, df, used_sheets, used_tables, prefix_sheet_names=prefix_sheet_names)

    return wb


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    wb = build_workbook()
    wb.save(OUTPUT_FILE)
    print(OUTPUT_FILE)


if __name__ == "__main__":
    main()
